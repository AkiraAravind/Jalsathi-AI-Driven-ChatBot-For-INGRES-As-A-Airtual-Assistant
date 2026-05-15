from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def add_page_number(paragraph):
    """Add a dynamic page number field in the footer paragraph."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    if level == 1:
        h.runs[0].font.size = Pt(16)
    return h


def add_bullet_list(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def main():
    base = Path(__file__).resolve().parent

    output_docx = base / "INGRES_UML_Documentation.docx"

    diagrams = [
        {
            "title": "Use Case Diagram",
            "file": "use case.jpeg",
            "purpose": "Shows who interacts with the system and what each role can do.",
            "details": [
                "Public User flow: register/login with email OTP, submit chat queries, and submit feedback.",
                "Government Official flow: access government portal, critical analytics, and active alert dashboards.",
                "Chatbot/System integration: Gemini-driven response generation with external email server used for OTP delivery.",
                "This diagram clarifies role-based access boundaries and core goals for each actor.",
            ],
        },
        {
            "title": "Class Diagram",
            "file": "class.jpeg",
            "purpose": "Represents the static structure of the chat platform and the relationships between classes/components.",
            "details": [
                "ChatRequest carries message and session identifiers from authenticated users.",
                "ChatEngine orchestrates embedding generation, vector query, prompt composition, and response synthesis.",
                "Gemini API is modeled as an external collaborator that returns generated responses.",
                "Database methods indicate specialized retrieval responsibilities (document, trend, and block risk search).",
                "ChatResponse formalizes answer text plus citations and language metadata.",
            ],
        },
        {
            "title": "Activity Diagram",
            "file": "Activity.jpeg",
            "purpose": "Describes the end-to-end control flow for handling user chat requests.",
            "details": [
                "Starts with user query submission and immediate authentication check.",
                "Unauthorized branch returns 401 and OTP re-entry path.",
                "Authorized branch converts query to embeddings and performs similarity search in Supabase pgvector.",
                "Retrieved context is appended into prompt before Gemini invocation.",
                "Final response is formatted with citations and returned to user.",
            ],
        },
        {
            "title": "State Chart Diagram",
            "file": "statechat.jpeg",
            "purpose": "Captures chat session lifecycle states and transitions.",
            "details": [
                "Lifecycle starts at Idle and transitions through New Chat Session and Sending OTP.",
                "After OTP verification, state moves to User Verified and then Processing User Query.",
                "System transitions to Context/AI Response and Chat Active for continuous conversation.",
                "When user ends chat, timer is stopped and session closes at Session Ended.",
            ],
        },
        {
            "title": "Sequence Diagram",
            "file": "sequence.jpeg",
            "purpose": "Shows time-ordered interactions across user, web app, backend API, Gemini API, and email server.",
            "details": [
                "Login lane: credential validation, OTP send, email dispatch, OTP entry, and login success.",
                "Chat lane: query submission to backend, processing, response generation, and citation-based answer return.",
                "Clarifies synchronous API calls and external dependency touchpoints.",
            ],
        },
        {
            "title": "Component Diagram",
            "file": "component.jpeg",
            "purpose": "Presents layered architecture and component responsibilities.",
            "details": [
                "User Layer includes Public Chat UI and Official Dashboard.",
                "Application Service Layer orchestrates chat and authentication services.",
                "Business logic centers around RAG Engine and Gemini Client.",
                "Data Layer includes Supabase vector database and SMTP email infrastructure.",
                "Dashed connectors indicate integration dependencies and pgvector retrieval path.",
            ],
        },
        {
            "title": "Deployment Diagram",
            "file": "deployment.jpeg",
            "purpose": "Maps runtime deployment nodes and communication paths.",
            "details": [
                "User and Admin browsers access a React frontend served from web server tier.",
                "Application server hosts Authentication Service, ChatEngine, and Feedback Service.",
                "Database server hosts Supabase (vector DB) and SMTP email server resources.",
                "Illustrates clean separation between presentation, application, and data tiers.",
            ],
        },
    ]

    code_snippet = Path(__file__).read_text(encoding="utf-8")

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("UML DIAGRAMS DOCUMENTATION")
    r.bold = True
    r.font.size = Pt(26)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("INGRES AI Command Center")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(16)

    s2 = doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s2.add_run("Comprehensive Analysis of System Design and Behavior")

    dline = doc.add_paragraph()
    dline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dline.add_run(f"Date: {date.today().strftime('%d %B %Y')}")

    doc.add_page_break()

    add_heading(doc, "Table of Contents", level=1)
    toc_items = [
        "1. Introduction",
        "2. Use Case Diagram",
        "3. Class Diagram",
        "4. Activity Diagram",
        "5. State Chart Diagram",
        "6. Sequence Diagram",
        "7. Component Diagram",
        "8. Deployment Diagram",
        "9. Implementation Mapping",
        "10. Appendix: Code Used to Generate This Document",
        "11. Conclusion",
    ]
    for item in toc_items:
        doc.add_paragraph(item)

    doc.add_page_break()

    add_heading(doc, "1. Introduction", level=1)
    doc.add_paragraph(
        "This report documents the full UML suite for the INGRES AI Command Center. "
        "The diagrams describe user interaction, static structure, runtime workflow, "
        "state transitions, architectural layers, and deployment topology. "
        "Together, they provide a unified engineering view of a secure RAG-based "
        "groundwater intelligence platform built with React, FastAPI, Supabase, and Gemini."
    )

    for idx, diagram in enumerate(diagrams, start=2):
        add_heading(doc, f"{idx}. {diagram['title']}", level=1)

        image_path = base / diagram["file"]
        if image_path.exists():
            pic = doc.add_picture(str(image_path), width=Inches(6.3))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph(f"Figure: {diagram['title']}")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.runs[0].italic = True
        else:
            miss = doc.add_paragraph(f"[Missing image: {diagram['file']}]")
            miss.runs[0].bold = True

        p = doc.add_paragraph()
        p.add_run("Purpose: ").bold = True
        p.add_run(diagram["purpose"])

        doc.add_paragraph("Explanation:")
        add_bullet_list(doc, diagram["details"])

    add_heading(doc, "9. Implementation Mapping", level=1)
    doc.add_paragraph(
        "The UML diagrams correspond to concrete implementation modules in this project:"
    )
    add_bullet_list(
        doc,
        [
            "Authentication and OTP flow maps to backend auth routes and SMTP integration.",
            "Chat processing maps to RAG Engine stages: embedding, retrieval, prompt assembly, and generation.",
            "Persistence and retrieval map to Supabase PostgreSQL + pgvector for semantic search.",
            "Frontend orchestration maps to React components/pages for login, dashboard, and chat workflows.",
        ],
    )

    add_heading(doc, "10. Appendix: Code Used to Generate This Document", level=1)
    doc.add_paragraph(
        "The following script was written and executed to generate this DOCX automatically."
    )

    code_lines = code_snippet.splitlines()
    for i, line in enumerate(code_lines, start=1):
        c = doc.add_paragraph(style="No Spacing")
        run = c.add_run(f"{i:03d}: {line}")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)

    add_heading(doc, "11. Conclusion", level=1)
    doc.add_paragraph(
        "The UML set demonstrates that INGRES is designed as a secure, layered, and extensible "
        "AI platform with clear actor boundaries, modular services, and well-defined data flow. "
        "This documentation can be used for architecture review, implementation alignment, and presentation."
    )

    # Footer page number
    section = doc.sections[0]
    footer = section.footer
    if not footer.paragraphs:
        footer.add_paragraph()
    add_page_number(footer.paragraphs[0])

    doc.save(output_docx)
    print(f"Created: {output_docx}")


if __name__ == "__main__":
    main()
