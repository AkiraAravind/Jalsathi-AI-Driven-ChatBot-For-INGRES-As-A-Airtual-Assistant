from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import os

# Create document
doc = Document()

# Add title
title = doc.add_heading('Results and Test Cases Documentation', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add introduction
intro = doc.add_paragraph()
intro.add_run('Project: INGRES - AI Integrated Hydrological Intelligence System\n').bold = True
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ==================== SECTION 1: WEBSITE SCREENSHOTS ====================
doc.add_heading('1. Website Screenshots', 1)

# Screenshots with descriptions
screenshots_info = [
    {
        'image': 'Entry-page.png',
        'title': 'Entry Page - Platform Landing Page',
        'description': '''Image Name: Entry-page.png

This is the main landing page of the INGRES platform introducing users to the Hydrological Intelligence system. The page displays a hero section with prominent branding and value proposition.

Visual Elements Shown:
• Header: "The Future of Hydrological Intelligence" - bold and centered
• Tagline: "Precision groundwater assessment using V-JEPA predictive architecture and sub-district level data mapping."
• Primary CTA Button: "Explore Dashboard" (blue button)
• Secondary Link: "About GEC Methodology"

Three Feature Cards Displayed:
1. Precision Mapping
   - Description: "Block-level sub-district data for all 36 States/UTs"
   - Icon: Compass/mapping icon

2. JEPA Prediction
   - Description: "V-JEPA predictive trajectory for at-risk districts."
   - Icon: Upward trend arrow

3. Manual Search
   - Description: "Instant retrieval of GEC-2015 methodology rules."
   - Icon: Search/query icon

Footer Information:
• "DEVELOPED BY TBP TEAM FOR SIH 2025/26"
• Partner organizations: Ministry of Jal Shakti, CGWB, IIT Hyderabad

Design Elements:
• Dark theme (navy/dark blue background)
• Modern gradient design
• Professional layout showcasing three core capabilities
• Clear navigation for first-time users

Purpose: This page serves as the entry point for all users, explaining INGRES's core functionality and encouraging them to explore the dashboard or learn more about the GEC methodology.'''
    },
    {
        'image': 'about-project.png',
        'title': 'About Project - Information Hub',
        'description': '''Image Name: about-project.png

This page provides comprehensive information about the INGRES project, its capabilities, and how different users can benefit from the platform.

Header Information:
• Project Badge: "CERTIFIED GEC-2015 INTELLIGENCE" (certification badge)
• Title: "Project INGRES" (in italics)
• Subtitle: "A Next-Generation Hydrological Intelligence System developed at Vasavi College of Engineering."

Four Information Sections (Cards):

1. Who can use this website?
   Icon: People/users icon (blue)
   Content: "This platform is built for: **Government Officials** (CGWB, SGWD) for executing advanced multi-year resource planning, and **General Citizens, Researchers, and Policymakers** who wish to explore national groundwater trends dynamically using natural language."

2. Why use the INGRES AI Hub?
   Icon: Lightbulb/star icon (yellow/orange)
   Content: "Because static PDFs and massive 150-column data tables are difficult to analyze. INGRES transforms raw data into instant, interactive visualizations and predictive trajectories allowing you to make rapid, data-driven decisions that could stabilize local water tables."

3. How is it helpful?
   Icon: Shield/security icon (blue)
   Content: "It leverages our cutting-edge **V-JEPA Engine** to predict which districts are trending toward 'Critical' or 'Over-Exploited' stages, catching risks **before** they occur. Furthermore, it strictly enforces GEC-2015 methodology guaranteeing scientific integrity in every query."

4. How to use
   Icon: Book/documentation icon (purple)
   Content: "• Navigate to the **AI ChatBot** console from the sidebar
   • Ask complex questions like: 'Compare the extraction stage of Hyderabad and Nalgonda in 2022'
   • Ask for visualizations by explicitly saying 'bar chart' or 'line plot'
   • Use Hindi or Telugu for localized answers naturally matching your query language."

Top Navigation:
• Logo: INGRES Team Jalsathi
• Menu Items: AI ChatBot, SIH Requirements, Gov Portal, About Project (highlighted), Support
• User Profile: Shows "Adhikari Sukla - GOVERNMENT OFFICIAL"

Purpose: Educational page explaining the platform's purpose, target users, and capabilities in detail.'''
    },
    {
        'image': 'user-login.png',
        'title': 'User Login - Authentication Portal',
        'description': '''Image Name: user-login.png

A centered authentication modal for users to connect to the INGRES dashboard. The interface presents a clean, focused login experience.

Modal Elements:

Header:
• Icon: Grid/menu icon (blue)
• Title: "AuthorizeAccess" (italicized, large font)
• Subtitle: "CONNECT TO INGRES DASHBOARD" (uppercase, gray text)

Login Form:
• Label: "OFFICIAL E-MAIL VECTOR"
• Input Field: Email address field pre-populated with "akiraarvind@outlook.com"
• Primary Action Button: "CONNECT HUB →" (blue, full-width)

Additional Options Below Button:
• "NEW RESIDENT" (left button/link for new user registration)
• "GOV PORTAL LOGIN" (right button/link for government user access)

Footer:
• "PROJECT INGRES // TEAM JALSATHI"

Design Characteristics:
• Dark background (navy/dark blue)
• Centered modal card design
• White text for contrast
• Blue accent colors for interactive elements
• Professional, clean aesthetic
• Email-based authentication approach

Purpose: Primary authentication point for regular users/residents to gain access to the INGRES AI ChatBot and dashboard features.'''
    },
    {
        'image': 'government-register.png',
        'title': 'Government User Registration (if visible)',
        'description': 'Image registration details not fully captured in available screenshot.'
    },
    {
        'image': 'goverenment-lohin.png',
        'title': 'Government User Login (if visible)',
        'description': 'Image not fully visible in current context.'
    },
    {
        'image': 'government-login-otp.png',
        'title': 'Government Login - OTP Verification (if visible)',
        'description': 'Image details not fully captured.'
    },
    {
        'image': 'government-login-otp-recived.jpeg',
        'title': 'Government Login - OTP Received (if visible)',
        'description': 'Image details not fully captured.'
    },
    {
        'image': 'government-portal-dashboard-upper-part1.png',
        'title': 'Government Portal Dashboard - Command Center',
        'description': '''Image Name: government-portal-dashboard-upper-part1.png

The Government Groundwater Command Portal provides government officials with a real-time overview of national groundwater status across all 36 states and union territories.

Header Section:
• Title: "Government Groundwater Command Portal"
• Subtitle: "NATIONAL WATER GOVERNANCE COMMAND GRID"
• Status: "Last Sync: 7:36:16 pm"
• Context: "Integrated decision console for officials to monitor groundwater extraction, compare all 36 states/UTs, and identify escalation zones through visual analytics and risk intelligence."

Key Performance Indicators (Upper Row):

1. NATIONAL STAGE OF EXTRACTION
   • Value: 0.00%
   • Trend: "+0.05% trending upward"
   • Indicates overall extraction percentage across nation

2. AT-RISK BLOCKS IDENTIFIED
   • Value: 0 (Yellow/orange highlight)
   • Note: "Requires Immediate ML verification"
   • Blocking indicator for critical zones

3. MONITORED JURISDICTIONS
   • Value: 36 (Blue highlight)
   • Note: "Fallback dataset active"
   • Represents all states/UTs under monitoring

Main Dashboard Section:

"Groundwater Level Graph - All 36 States/UTs"
• Subtitle: "Stage of extraction (%) with bar + line trend overlay for official comparison"
• Chart Type: Bar chart with line overlay
• X-axis: Shows different states/regions
• Y-axis: Percentage scale (0-140)
• Colors: Green, yellow, and red bars representing different extraction levels
• Notable spikes visible in central data
• Button: "FALLBACK ANALYTICAL VIEW" (yellow)

Operational Features Panel (Right Side):
1. Aquifer Stress Scan
   - "Automated extraction stress scoring for every jurisdiction with threshold-based flags."

2. Real-time Risk Signals
   - "Live telemetry pulse aligned with official reporting windows and policy checkpoints."

3. 36-State Coverage
   - "State and UT level monitoring view with a single comparative governance pane."

4. Priority Escalation
   - "Critical-state queue for immediate intervention and district-level driftdown."

Purpose: Provides government officials with comprehensive real-time monitoring and risk assessment of national groundwater resources.'''
    },
    {
        'image': 'government-portal-dashboard-upper-part2.png',
        'title': 'Government Portal Dashboard - Analytics & Chat Interface',
        'description': '''Image Name: government-portal-dashboard-upper-part2.png

The dashboard continuation shows the AI ChatBot integration within the government portal, demonstrating query capabilities and data visualization features.

Left Sidebar:
• Section: "RECENT INTELLIGENCE"
• Navigation Item: "New Exploration" (with folder icon)
• Session History showing:
  - Multiple "Session sessio..." entries with dates (e.g., 2025-04-10)
  - Collapsible history of recent queries

Main Content Area:

Chat Interface Header:
• Database: "INGRES AI Core"
• Status: "HYDROLOGICAL LINK ACTIVE" (green indicator)
• Features: PDF (export), MULTI-LINGUAL (language support buttons)

Recent Query Example:
• User Question: "What was the Stage of Groundwater Extraction in Nalgonda district in 2022?"
• System Response: Detailed answer with data sources

Response Content Includes:
• Direct Answer: "In 2022, the Stage of Groundwater Extraction in Nalgonda district, Telangana, was 41.3%, categorizing it as Safe [SOURCE: INGRES Assessment Data, TELANGANA — Nalgonda District, Year 2022]."

• Additional Details Provided:
  - AEGR (Annual Extractable GW Resource): 138890.26 Ham
  - Total Extraction: 57418.04 Ham
  - Net GW Availability for Future Use: 83388.72 Ham
  - Groundwater Trend (IEPA): Stable with -0.3% per year change

• Query Context Formula:
  - "Stage of Extraction (%) = (Total Annual GW Extraction / AEGR) × 100"

• Sources Listed:
  - INGRES Assessment Data, TELANGANA — Multiple districts and years
  - GEC User Manual, Section 2.5

• Source Citations: Multiple data source tags at bottom showing specific datasets used

Chat Input Area:
• Placeholder: "Ask about district trends, GEC formulas, or categorization..."
• Allows follow-up questions and continuous conversation

Purpose: Demonstrates the AI ChatBot's capability to provide detailed, sourced answers to government queries about hydrological data with multi-lingual support and data visualization.'''
    },
    {
        'image': 'feedback.png',
        'title': 'Feedback Form (if visible)',
        'description': 'Image details to be captured.'
    },
    {
        'image': 'feedback-form-recived.jpeg',
        'title': 'Feedback Confirmation (if visible)',
        'description': 'Image details to be captured.'
    },
]

# Add screenshots with descriptions
for idx, item in enumerate(screenshots_info, 1):
    image_path = f'e:\\INGRES_TBP\\outputs_docmentation\\{item["image"]}'
    
    if os.path.exists(image_path):
        # Add heading for each screenshot
        heading = doc.add_heading(f'{idx}. {item["title"]}', 2)
        
        # Add image
        try:
            doc.add_picture(image_path, width=Inches(5.5))
            last_paragraph = doc.paragraphs[-1]
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Add image name caption
            caption = doc.add_paragraph()
            caption_run = caption.add_run(f"Figure {idx}: {item['image']}")
            caption_run.font.size = Pt(9)
            caption_run.font.italic = True
            caption_run.font.color.rgb = RGBColor(128, 128, 128)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            doc.add_paragraph(f'[Error loading image: {item["image"]}]')
        
        # Add description
        desc_para = doc.add_paragraph(item['description'])
        desc_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Add spacing
        doc.add_paragraph()
    else:
        print(f"Warning: Image not found - {image_path}")

# Save document
output_path = 'e:\\INGRES_TBP\\Results_and_TestCases_Documentation_UPDATED.docx'
doc.save(output_path)
print(f"Accurate Document created successfully: {output_path}")
print("\nNote: Please add more images and their accurate descriptions to complete the document.")
